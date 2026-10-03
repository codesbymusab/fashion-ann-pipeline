import os
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_academic_docx():
    doc = Document()

    # Page Margins - Standard Academic (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = False

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("CS4069 MLOps | Assignment 3: Git + DVC + Google Drive Pipeline")
        hrun.font.name = "Times New Roman"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(128, 128, 128)

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("National University of Computer & Emerging Sciences (FAST-NUCES) Lahore")
        frun.font.name = "Times New Roman"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(128, 128, 128)

    # Document Styles
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(11)
    normal_font.color.rgb = RGBColor(30, 30, 30)

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("End-to-End Machine Learning Versioning and Pipeline Orchestration Using Git, DVC, and Cloud Remote Storage")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(17)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(16, 44, 87)

    # Subtitle / Course
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run("Course Assignment 3 Technical Report | Machine Learning Operations (MLOps)")
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(70, 70, 70)

    # Author Card
    author_table = doc.add_table(rows=1, cols=1)
    author_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = author_table.cell(0, 0)
    set_cell_background(cell, "F2F4F7")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    ap = cell.paragraphs[0]
    ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ar1 = ap.add_run("Muhammad Musab (Roll No: 23L-0806)\n")
    ar1.font.bold = True
    ar1.font.size = Pt(11.5)
    ar1.font.color.rgb = RGBColor(16, 44, 87)

    ar2 = ap.add_run("Department of Computer Science, FAST-NUCES Lahore, Pakistan\n")
    ar2.font.size = Pt(10)
    ar3 = ap.add_run("Institutional Email: l230806@lhr.nu.edu.pk | GitHub Remote: codesbymusab/fashion-ann-pipeline")
    ar3.font.size = Pt(9.5)
    ar3.font.italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Abstract Box
    abs_table = doc.add_table(rows=1, cols=1)
    abs_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    abs_cell = abs_table.cell(0, 0)
    set_cell_background(abs_cell, "EBF3FA")
    set_cell_margins(abs_cell, top=140, bottom=140, left=200, right=200)
    ab_p = abs_cell.paragraphs[0]
    ab_run_title = ab_p.add_run("Abstract—")
    ab_run_title.font.bold = True
    ab_run_title.font.size = Pt(10)
    ab_run_title.font.color.rgb = RGBColor(16, 44, 87)
    ab_text = ab_p.add_run(
        "Production machine learning workflows present dual versioning requirements: lightweight source code must remain trackable in a distributed version control system, while multi-megabyte data arrays and trained neural network binaries must be managed through content-addressable storage without bloat. This report presents the full implementation of an end-to-end MLOps pipeline on the Fashion-MNIST benchmark utilizing Git for source versioning, Data Version Control (DVC) for artifact tracking, and Google Drive as the remote object store. We evaluate advanced Git operations including stash lifecycles, atomic renames, non-fast-forward hotfix integration, and rebase linearization. Subsequently, we construct a 4-stage modular pipeline governed by params.yaml and dvc.yaml, establishing automated reproducibility and intelligent stage-skipping across parameter variations (dense units: 128 vs 256). Finally, we simulate a distributed team collaboration scenario, demonstrating the concurrent emergence of a textual code conflict in data preprocessing and a hash collision in dvc.lock, followed by a formal authoritative reconciliation methodology. All results, metrics diffs, and verification audits are documented."
    )
    ab_text.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(16, 44, 87)
        return h

    def add_heading_2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(3)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(40, 80, 130)
        return h

    def add_figure(img_path, caption):
        if Path(img_path).exists():
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(6)
            p_img.paragraph_format.space_after = Pt(2)
            run = p_img.add_run()
            run.add_picture(str(img_path), width=Inches(6.2))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(8)
            r_cap = p_cap.add_run(caption)
            r_cap.font.name = "Times New Roman"
            r_cap.font.size = Pt(9)
            r_cap.font.italic = True
            r_cap.font.color.rgb = RGBColor(90, 90, 90)

    # ----------------- SECTION 1 -----------------
    add_heading_1("1. Introduction & Theoretical Background")
    p = doc.add_paragraph(
        "Machine Learning Operations (MLOps) addresses the fundamental friction between standard software versioning and data science experimentation. Traditional source code version control (e.g., Git) is designed for delta-compression of UTF-8 text. Tracking large, opaque binary artifacts—such as raw dataset matrices, preprocessed tensors, and serialized model weights—directly in Git repositories causes exponential repository bloat, degrades cloning bandwidth, and makes historical traversal impractical."
    )
    p = doc.add_paragraph(
        "Data Version Control (DVC) solves this architectural challenge by introducing lightweight pointer files (content-addressable md5 hashes) into Git, while delegating the underlying payloads to remote cloud storage. In this assignment, we implement a production-grade pipeline for the Fashion-MNIST classification benchmark, examining the complete lifecycle from repository initialization to multi-developer merge conflict resolution."
    )

    # ----------------- SECTION 2 -----------------
    add_heading_1("2. Part A: Advanced Git Operations & Workflow Mechanics")
    p = doc.add_paragraph(
        "Part A focuses on rigorous version control hygiene, branch topology management, and defensive repository operations across tasks A1 through A8."
    )

    add_heading_2("2.1 Repository Setup and Initial Branching (Tasks A1 & A2)")
    p = doc.add_paragraph(
        "The repository was initialized locally on the main branch. A deliberate typo ('# Fashion-MNIST ANN Pipelnie') was introduced in README.md to simulate an upstream production defect for subsequent hotfixing. A strict .gitignore was authored to ensure Python virtual environments, compiled bytecode, and HDF5 weight files (*.h5) are excluded from direct Git tracking. An initial commit was recorded, remote origin was bound to GitHub, and the dev branch was established for staging iterative script development across six distinct commits."
    )

    add_heading_2("2.2 Commit Log Topology Analysis (Task A3)")
    p = doc.add_paragraph(
        "Commit history inspection provides insight into branching topology and code evolution. Four complementary log variants were evaluated as shown in Figure 1."
    )
    add_figure("screenshots/A3_git_log_variants.png", "Figure 1: Git commit history graph (--graph --all) and unmerged commit filtering (main..dev).")

    # Table 1: Log explanations
    table1 = doc.add_table(rows=5, cols=2)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Command", "Analytical Significance & Architectural Purpose"]
    for i, h in enumerate(headers):
        cell = table1.cell(0, i)
        cell.paragraphs[0].text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9.5)
        set_cell_background(cell, "EAECEF")
        set_cell_margins(cell, 80, 80, 100, 100)

    log_data = [
        ("git log --oneline --graph --all", "Visualizes the full multi-branch topology, depicting divergence points, active branch pointers, and merge nodes across all local and remote references."),
        ("git log --stat -3", "Provides high-level churn metrics per commit (files modified, insertions, deletions) without cluttering the output with full syntactic diffs."),
        ("git log -p -1", "Generates the complete line-by-line patch delta for the head commit, essential for granular code reviews and debugging subtle regressions."),
        ("git log main..dev", "Isolates only unmerged commits reachable from dev that have not yet been integrated into main, serving as a pre-merge inspection gate.")
    ]
    for row_idx, (cmd, desc) in enumerate(log_data, start=1):
        c0 = table1.cell(row_idx, 0)
        c0.paragraphs[0].text = cmd
        c0.paragraphs[0].runs[0].font.size = Pt(8.5)
        c0.paragraphs[0].runs[0].font.name = "Consolas"
        c1 = table1.cell(row_idx, 1)
        c1.paragraphs[0].text = desc
        c1.paragraphs[0].runs[0].font.size = Pt(9)
        set_cell_margins(c0, 60, 60, 80, 80)
        set_cell_margins(c1, 60, 60, 80, 80)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_heading_2("2.3 Differential Code Analysis: Two-Dot vs. Three-Dot Syntax (Task A4)")
    p = doc.add_paragraph(
        "To inspect changes across staging levels and branch tips, differential analysis commands were executed (Figure 2). A critical distinction exists between two-dot and three-dot branch comparisons:"
    )
    p_b1 = doc.add_paragraph()
    p_b1.paragraph_format.left_indent = Inches(0.25)
    r1 = p_b1.add_run("1. Two-Dot Comparison (git diff main..dev): ")
    r1.font.bold = True
    p_b1.add_run(
        "This compares the exact tips of both branches. It displays all net differences, meaning any commit present exclusively on main (such as an independent hotfix) will appear as a reverse/negative deletion in dev's context."
    )
    p_b2 = doc.add_paragraph()
    p_b2.paragraph_format.left_indent = Inches(0.25)
    r2 = p_b2.add_run("2. Three-Dot Comparison (git diff main...dev): ")
    r2.font.bold = True
    p_b2.add_run(
        "This computes the symmetric merge-base (the most recent common ancestor commit) and compares dev strictly against that ancestor. It isolates exclusively the work authored on dev, filtering out all concurrent divergence on main."
    )
    add_figure("screenshots/A4_git_diff_variants.png", "Figure 2: Verification of unstaged changes, staged diffs, and relative branch diffs.")

    add_heading_2("2.4 Work-in-Progress Isolation via Git Stash (Task A5)")
    p = doc.add_paragraph(
        "During active development on dev, an uncommitted debug line was added to src/preprocess.py. An immediate attempt to switch branches via git checkout main was halted by Git's index-protection mechanism, which prevents dirty working directories from overwriting untracked or divergent branch files. To circumvent this without creating premature commits, the working state was pushed to the stash stack (git stash push -m 'wip: debug print'). Following checkout to main and return to dev, the state was re-applied and cleared using git stash pop (Figure 3)."
    )
    add_figure("screenshots/A5_git_stash.png", "Figure 3: Git stash sequence: checkout rejection, dirty state stashing, and clean stack pop.")

    add_heading_2("2.5 Hotfix Branching, Non-Fast-Forward Merge & Rebase Conflict (Task A6)")
    p = doc.add_paragraph(
        "To remedy the README typo on production, branch hotfix/readme-typo was spawned from main. The correction was committed and integrated into main via a non-fast-forward merge (git merge --no-ff), preserving the explicit branch boundary in graph history. When dev was subsequently rebased onto main, a content conflict halted execution at line 1 of README.md. The conflicting markers were manually inspected, reconciled to retain both the spelling fix and feature addition, staged, and finalized via git rebase --continue, producing a unified linear history (Figure 4)."
    )
    add_figure("screenshots/A6_rebase_conflict.png", "Figure 4: Upstream hotfix merge, rebase collision on README.md, and conflict resolution.")

    add_heading_2("2.6 Reset Mechanics: Soft vs. Hard Recovery (Task A7)")
    p = doc.add_paragraph(
        "On an isolated scratch branch, we evaluated the operational difference between soft and hard resets (Figure 5):"
    )
    p_b3 = doc.add_paragraph()
    p_b3.paragraph_format.left_indent = Inches(0.25)
    r3 = p_b3.add_run("• Soft Reset (git reset --soft HEAD~1): ")
    r3.font.bold = True
    p_b3.add_run("Moves the branch HEAD pointer backwards one commit, leaving the staging area (index) and the working tree untouched. All changes remain staged and preserved on disk, allowing commit squash or message amendment.")
    p_b4 = doc.add_paragraph()
    p_b4.paragraph_format.left_indent = Inches(0.25)
    r4 = p_b4.add_run("• Hard Reset (git reset --hard HEAD~1): ")
    r4.font.bold = True
    p_b4.add_run("Moves HEAD backwards while simultaneously synchronizing both the index and working tree to match the target commit. All uncommitted and staged modifications are permanently expunged from the disk.")
    add_figure("screenshots/A7_git_reset.png", "Figure 5: Empirical observation of soft reset (staged preservation) vs hard reset (disk deletion).")

    add_heading_2("2.7 Atomic File Operations and Rename Tracking (Task A8)")
    p = doc.add_paragraph(
        "To organize loose repository scripts, check_data.py was relocated into src/check_data.py using git mv, while scratch_notes.txt was expunged via git rm. Executing git show --stat -M HEAD verified that Git's internal similarity index detected a 100% rename event rather than treating the file as an unrelated deletion and insertion pair (Figure 6)."
    )
    add_figure("screenshots/A8_git_mv_rm.png", "Figure 6: Git tree tracking showing 100% rename detection and clean file removal.")

    # ----------------- SECTION 3 -----------------
    add_heading_1("3. Part B: Modular Machine Learning Pipeline Architecture")
    p = doc.add_paragraph(
        "The project implements a decoupled 4-stage pipeline for Fashion-MNIST image classification (60,000 training images, 10,000 test images across 10 clothing categories). Each component executes as an independent script driven by central configurations:"
    )
    p_s1 = doc.add_paragraph()
    p_s1.paragraph_format.left_indent = Inches(0.25)
    p_s1.add_run("1. Stage 1: Ingestion (src/prepare.py): ").font.bold = True
    p_s1.add_run("Downloads Fashion-MNIST from keras.datasets and writes raw uncompressed arrays (x_train.npy, y_train.npy, x_test.npy, y_test.npy) into data/raw/. Using plain .npy preserves identical byte layouts across invocations to stabilize cryptographic hashing.")

    p_s2 = doc.add_paragraph()
    p_s2.paragraph_format.left_indent = Inches(0.25)
    p_s2.add_run("2. Stage 2: Preprocessing (src/preprocess.py): ").font.bold = True
    p_s2.add_run("Normalizes uint8 pixels [0, 255] into float32 [0.0, 1.0], executes a 10% stratified train/validation split using params.yaml (test_size=0.1, seed=42), and outputs 6 array files to data/processed/.")

    p_s3 = doc.add_paragraph()
    p_s3.paragraph_format.left_indent = Inches(0.25)
    p_s3.add_run("3. Stage 3: Model Construction & Training (src/train.py): ").font.bold = True
    p_s3.add_run("Builds a sequential Artificial Neural Network (Flatten 28x28 -> Dense(ReLU) -> Dropout(0.2) -> Dense(10, Softmax)). Compiled with Adam optimizer and sparse categorical crossentropy. Logs epoch training history to models/history.csv and serializes weights to models/model.h5.")

    p_s4 = doc.add_paragraph()
    p_s4.paragraph_format.left_indent = Inches(0.25)
    p_s4.add_run("4. Stage 4: Evaluation & Reporting (src/evaluate.py): ").font.bold = True
    p_s4.add_run("Loads models/model.h5, evaluates test loss and accuracy, serializes metrics to metrics.json, and generates confusion_matrix.png using scikit-learn ConfusionMatrixDisplay.")

    p_perf = doc.add_paragraph(
        "Baseline training achieved convergence within 10 epochs: Training Accuracy: 89.36%, Training Loss: 0.2846, Validation Accuracy: 89.02%, Validation Loss: 0.3064."
    )

    # ----------------- SECTION 4 -----------------
    add_heading_1("4. Part C: DVC Remote Storage Integration with Google Drive")
    p = doc.add_paragraph(
        "Data Version Control was initialized on the dev branch. A designated remote store named gdrive_storage was configured targeting Google Drive folder ID 1FdsWzNpKhXq4zAcrmCMqHFjtrnoguXLS."
    )

    add_heading_2("4.1 OAuth Security Hardening & Credential Isolation")
    p = doc.add_paragraph(
        "Public default OAuth credentials bundled with PyDrive2 often encounter Google quota throttling and security blocks. To guarantee enterprise reliability, a custom Google Cloud Project was provisioned, enabling the Google Drive API and registering a Desktop Application OAuth client. Crucially, credentials were split into two security tiers:"
    )
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.left_indent = Inches(0.25)
    p_sec.add_run("• Public Metadata (.dvc/config): ").font.bold = True
    p_sec.add_run("Stores the non-sensitive client ID and remote URL, safely committed to Git.\n")
    p_sec.add_run("• Secret Key (.dvc/config.local): ").font.bold = True
    p_sec.add_run("Stores the sensitive client secret using the --local flag. Verified by git check-ignore -v .dvc/config.local, keeping all secrets strictly isolated from version control.")

    p_audit = doc.add_paragraph(
        "An automated audit via git ls-files | grep -iE 'cred|token|secret' yielded 0 tracked credentials. Executing dvc push successfully uploaded all 14 artifact files (~275 MB total) to Google Drive in content-addressed hash format (Figure 7)."
    )
    add_figure("screenshots/BC_dvc_setup_push.png", "Figure 7: DVC remote configuration, local secret isolation, security audit, and Drive sync.")

    # ----------------- SECTION 5 -----------------
    add_heading_1("5. Part D: DVC Pipeline Automation & Hyperparameter Tuning")
    p = doc.add_paragraph(
        "In Part D, artifact ownership was transitioned from static .dvc tracking files to an executable pipeline defined in dvc.yaml. Stages explicitly define dependencies (deps), output directories (outs), tracked parameters (params), and metrics."
    )

    add_heading_2("5.1 Baseline Execution (Version 1)")
    p = doc.add_paragraph(
        "Executing dvc repro for the baseline (dense_units: 128) processed the full 4-stage pipeline, establishing the state lock in dvc.lock. Baseline test evaluation yielded: test_accuracy = 0.8666 and test_loss = 0.36053. The state was committed to Git and tagged as v1."
    )

    add_heading_2("5.2 Hyperparameter Variation & Intelligent Stage Skipping (Version 2)")
    p = doc.add_paragraph(
        "To evaluate pipeline automation, params.yaml was edited: train.dense_units was scaled from 128 to 256. Running dvc status immediately flagged stage train as out-of-date due to modified parameter dependencies. Re-executing dvc repro demonstrated intelligent caching:"
    )
    p_skip = doc.add_paragraph()
    p_skip.paragraph_format.left_indent = Inches(0.25)
    p_skip.add_run("• Stage 'prepare': Skipped (inputs and code unchanged).\n")
    p_skip.add_run("• Stage 'preprocess': Skipped (raw arrays and preprocess params unchanged).\n")
    p_skip.add_run("• Stage 'train': Re-executed (detected modified parameter train.dense_units).\n")
    p_skip.add_run("• Stage 'evaluate': Re-executed (triggered by updated models/model.h5 output).")

    p = doc.add_paragraph(
        "This smart skipping saved considerable computation and eliminated redundant dataset downloads. The v2 experiment was tagged and pushed to remote storage (Figure 8)."
    )
    add_figure("screenshots/D_dvc_pipeline_v1_v2.png", "Figure 8: Reproduction of v1 baseline, param diff, selective stage skipping, and metrics comparison.")

    # Table 2: v1 vs v2
    table2 = doc.add_table(rows=4, cols=4)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t2_headers = ["Metric / Parameter", "v1 Baseline (128 Units)", "v2 Experiment (256 Units)", "Observed Impact"]
    for i, h in enumerate(t2_headers):
        c = table2.cell(0, i)
        c.paragraphs[0].text = h
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.size = Pt(9.5)
        set_cell_background(c, "EAECEF")
        set_cell_margins(c, 80, 80, 100, 100)

    exp_data = [
        ("train.dense_units", "128", "256", "+128 hidden units (2x capacity)"),
        ("test_accuracy", "0.8666 (86.66%)", "0.8774 (87.74%)", "+0.0108 (+1.08% improvement)"),
        ("test_loss", "0.36053", "0.34467", "-0.01586 (Improved generalization)")
    ]
    for row_idx, row in enumerate(exp_data, start=1):
        for col_idx, val in enumerate(row):
            c = table2.cell(row_idx, col_idx)
            c.paragraphs[0].text = val
            c.paragraphs[0].runs[0].font.size = Pt(9)
            if col_idx == 0:
                c.paragraphs[0].runs[0].font.name = "Consolas"
            set_cell_margins(c, 60, 60, 80, 80)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ----------------- SECTION 6 -----------------
    add_heading_1("6. Part E: Multi-Developer Collaboration & Dual Conflict Resolution")
    p = doc.add_paragraph(
        "A realistic software engineering team scenario was simulated by branching two concurrent contributors from main to investigate preprocessing enhancements."
    )

    add_heading_2("6.1 Divergent Feature Engineering Implementations")
    p_dev1 = doc.add_paragraph()
    p_dev1.paragraph_format.left_indent = Inches(0.25)
    p_dev1.add_run("1. Contributor 1 (Branch teammate-sim): ").font.bold = True
    p_dev1.add_run("Implemented full dataset standardization using Fashion-MNIST empirical constants: (x / 255.0 - 0.2860) / 0.3530. Regenerated processed data via dvc repro preprocess, pushed artifacts to Drive, and committed to teammate-sim.")

    p_dev2 = doc.add_paragraph()
    p_dev2.paragraph_format.left_indent = Inches(0.25)
    p_dev2.add_run("2. Contributor 2 (Branch main): ").font.bold = True
    p_dev2.add_run("Implemented zero-centered normalization: (x / 255.0) - 0.5. Regenerated processed data via dvc repro preprocess, pushed artifacts to Drive, and committed to main.")

    add_heading_2("6.2 Emergence of Dual Code & Data Pointer Conflicts")
    p = doc.add_paragraph(
        "When attempting git merge teammate-sim into main, Git encountered two simultaneous collisions:\n"
        "1. Source Code Conflict (src/preprocess.py): Both branches modified identical lines inside the normalize() function.\n"
        "2. Data Pointer Collision (dvc.lock): In a pipeline-driven project, data hashes for data/processed reside within dvc.lock. Since both branches regenerated data from different equations, the md5 directory hashes collided."
    )

    add_heading_2("6.3 Authoritative Reconciliation Workflow")
    p = doc.add_paragraph(
        "Resolving this required a principled two-stage strategy (Figure 9):\n"
        "1. Code Reconciliation: We unified the engineering rationale. Both developers intended zero-centered distribution; standardization is the mathematically superior formulation. We authored a unified normalize() scaling to [0, 1] followed by mean/std standardization.\n"
        "2. Data Pointer Clearing & Regeneration: Because the merged code produces a 3rd distinct version of the dataset that neither branch previously created, conflict markers were cleared using git checkout --ours dvc.lock, synced via dvc checkout, and regenerated from source using dvc repro."
    )
    p = doc.add_paragraph(
        "The reconciled pipeline executed cleanly: test_accuracy = 0.8715. Running dvc repro immediately afterwards printed 'didn't change, skipping' across all 4 stages, verifying 100% deterministic reproducibility."
    )
    add_figure("screenshots/E_merge_conflict_resolution.png", "Figure 9: Merge collision in Python code and dvc.lock, reconciliation, and data regeneration.")

    # ----------------- SECTION 7 -----------------
    add_heading_1("7. Final Verification & Topology Graph")
    p = doc.add_paragraph(
        "All code branches (main, dev, teammate-sim) and tags (v1, v2) were pushed to GitHub. The complete unsquashed commit topology is depicted in Figure 10."
    )
    add_figure("screenshots/E_final_git_graph.png", "Figure 10: Complete commit graph topology showing hotfix merge, dev rebase, and teammate merge.")

    # Verification Table
    v_table = doc.add_table(rows=6, cols=3)
    v_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    v_headers = ["Audit Requirement", "Verification Command", "Observed Outcome"]
    for i, h in enumerate(v_headers):
        c = v_table.cell(0, i)
        c.paragraphs[0].text = h
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.size = Pt(9.5)
        set_cell_background(c, "EAECEF")
        set_cell_margins(c, 80, 80, 100, 100)

    v_data = [
        ("Complete Unsquashed Git History", "git log --oneline --graph --all", "Verified: Clean topology showing hotfix, rebase, and merge."),
        ("Tagging & Release Provenance", "git tag -l && git push origin --tags", "Verified: v1 (0.8666) and v2 (0.8774) pushed to remote."),
        ("DVC Pipeline Lock Integrity", "dvc repro (idempotency check)", "Verified: All stages report 'didn't change, skipping'."),
        ("Cloud Remote Cache Synchronization", "dvc status -c", "Verified: 'Cache and remote gdrive_storage are in sync'."),
        ("Security Audit: Zero Credentials in Git", "git ls-files | grep -iE 'cred|token|secret'", "Verified: 0 matching files; secret kept in .dvc/config.local.")
    ]
    for row_idx, row in enumerate(v_data, start=1):
        for col_idx, val in enumerate(row):
            c = v_table.cell(row_idx, col_idx)
            c.paragraphs[0].text = val
            c.paragraphs[0].runs[0].font.size = Pt(8.5)
            if col_idx == 1:
                c.paragraphs[0].runs[0].font.name = "Consolas"
            set_cell_margins(c, 60, 60, 80, 80)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ----------------- SECTION 8 -----------------
    add_heading_1("8. Conclusion")
    p = doc.add_paragraph(
        "This project successfully operationalized the full spectrum of Git and DVC integration. By maintaining code changes in Git and large multidimensional arrays in Google Drive, the pipeline achieves high speed, reproducibility, and auditability. The demonstrated resolution of dual code and DVC lock conflicts equips engineering teams with a robust methodology for collaborative data science and continuous model integration."
    )

    out_file = "Assignment3_Academic_Report.docx"
    doc.save(out_file)
    print(f"Successfully generated academic Word document: {out_file}")

if __name__ == "__main__":
    create_academic_docx()
